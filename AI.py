import sys
import os
import json
import tempfile
import docx
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QTextEdit, QFileDialog, QMessageBox,
    QVBoxLayout, QWidget, QMenu, QScrollArea, QPushButton, QInputDialog,
    QLabel, QLineEdit, QDialog, QGridLayout, QMenuBar, QWidgetAction
)
from PyQt6.QtGui import QAction
from PyQt6.QtCore import QThread, pyqtSignal
from fpdf import FPDF
from deep_translator import GoogleTranslator
import sounddevice as sd
import queue
import vosk
import json as js
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.backends import default_backend
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
import base64
import secrets

# -------- Voice Thread -------- #
class VoiceTypingThread(QThread):
    text_output = pyqtSignal(str)

    def __init__(self, model_path):
        super().__init__()
        self.model_path = model_path
        self.running = True

    def run(self):
        q = queue.Queue()

        if not os.path.exists(self.model_path):
            self.text_output.emit("[Error: Vosk model not found at path]")
            return

        try:
            model = vosk.Model(self.model_path)
        except Exception as e:
            self.text_output.emit(f"[Error loading Vosk model: {e}]")
            return

        device = None
        try:
            samplerate = int(sd.query_devices(device, 'input')['default_samplerate'])
        except Exception as e:
            self.text_output.emit(f"[Error accessing microphone: {e}]")
            return

        def callback(indata, frames, time, status):
            q.put(bytes(indata))

        with sd.RawInputStream(samplerate=samplerate, blocksize=8000, dtype='int16',
                               channels=1, callback=callback):
            rec = vosk.KaldiRecognizer(model, samplerate)
            while self.running:
                data = q.get()
                if rec.AcceptWaveform(data):
                    result = js.loads(rec.Result())
                    if result.get("text"):
                        self.text_output.emit(result["text"])

    def stop(self):
        self.running = False


# -------- Encryption Functions -------- #
def generate_key(password: str, salt: bytes) -> bytes:
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(),
        length=32,
        salt=salt,
        iterations=100000,
        backend=default_backend()
    )
    return kdf.derive(password.encode())


def encrypt_data(data: bytes, password: str) -> bytes:
    salt = os.urandom(16)
    key = generate_key(password, salt)
    iv = os.urandom(16)
    cipher = Cipher(algorithms.AES(key), modes.CFB(iv), backend=default_backend())
    encryptor = cipher.encryptor()
    encrypted = encryptor.update(data) + encryptor.finalize()
    return salt + iv + encrypted


def decrypt_data(encrypted_data: bytes, password: str) -> bytes:
    salt = encrypted_data[:16]
    iv = encrypted_data[16:32]
    data = encrypted_data[32:]
    key = generate_key(password, salt)
    cipher = Cipher(algorithms.AES(key), modes.CFB(iv), backend=default_backend())
    decryptor = cipher.decryptor()
    return decryptor.update(data) + decryptor.finalize()


# -------- Main App -------- #
class TextCraftAI(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("TextCraft AI")
        self.setGeometry(100, 100, 1000, 700)
        self.text_edit = QTextEdit()
        self.setCentralWidget(self.text_edit)
        self.model_path = "C:/vosk-model-small-en-us-0.15"  # Change if needed
        self.voice_thread = None
        self.current_password = None
        self.recovery_key = None
        self.create_menu()

    def create_menu(self):
        menu_bar = self.menuBar()

        # File Menu
        file_menu = menu_bar.addMenu("File")
        file_menu.addAction("Open", self.open_file)
        file_menu.addAction("Save", self.save_file)
        file_menu.addAction("Save as PDF", self.save_pdf)
        file_menu.addAction("Save as DOC", self.save_doc)
        file_menu.addAction("Save History", self.save_history)

        # AI Menu
        ai_menu = menu_bar.addMenu("AI")
        ai_menu.addAction("Generate with AI", self.generate_ai)
        ai_menu.addAction("Summary with AI", self.summarize_ai)

        # Translator Menu
        translator_menu = menu_bar.addMenu("Translator")
        scroll_area = QScrollArea()
        scroll_widget = QWidget()
        grid_layout = QGridLayout(scroll_widget)
        languages = ["en", "es", "fr", "de", "hi", "bn", "te", "ta", "kn", "ml",
                     "gu", "mr", "ur", "pa", "or", "as", "ne", "si", "zh", "ja",
                     "ru", "it", "pt", "ar", "ko", "tr", "fa", "pl", "uk", "ro",
                     "nl", "sv"]
        for i, lang in enumerate(languages):
            button = QPushButton(lang.upper())
            button.clicked.connect(lambda checked, l=lang: self.translate_text(l))
            grid_layout.addWidget(button, i // 4, i % 4)
        scroll_area.setWidgetResizable(True)
        scroll_area.setWidget(scroll_widget)
        translator_action = QWidgetAction(self)
        translator_action.setDefaultWidget(scroll_area)
        translator_menu.addAction(translator_action)

        # Voice Menu
        voice_menu = menu_bar.addMenu("Voice")
        voice_menu.addAction("Start", self.start_voice_typing)
        voice_menu.addAction("Stop", self.stop_voice_typing)

        # Password Menu
        pass_menu = menu_bar.addMenu("Password")
        pass_menu.addAction("Set Password", self.set_password)
        pass_menu.addAction("Save Encrypted", self.save_encrypted_file)
        pass_menu.addAction("Open Encrypted", self.open_encrypted_file)

    # File Functions
    def open_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open File", "", "Text Files (*.txt)")
        if path:
            with open(path, 'r', encoding='utf-8') as f:
                self.text_edit.setText(f.read())

    def save_file(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save File", "", "Text Files (*.txt)")
        if path:
            with open(path, 'w', encoding='utf-8') as f:
                f.write(self.text_edit.toPlainText())

    def save_pdf(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save PDF", "", "PDF Files (*.pdf)")
        if path:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            for line in self.text_edit.toPlainText().split('\n'):
                pdf.cell(200, 10, txt=line, ln=True)
            pdf.output(path)

    def save_doc(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save DOC", "", "Word Files (*.docx)")
        if path:
            doc = docx.Document()
            doc.add_paragraph(self.text_edit.toPlainText())
            doc.save(path)

    def save_history(self):
        with open("history.txt", "a", encoding='utf-8') as f:
            f.write(self.text_edit.toPlainText() + "\n---\n")

    # AI Placeholder
    def generate_ai(self):
        QMessageBox.information(self, "AI", "Generate with AI - Not implemented")

    def summarize_ai(self):
        QMessageBox.information(self, "AI", "Summary with AI - Not implemented")

    # Translator
    def translate_text(self, lang):
        text = self.text_edit.toPlainText()
        try:
            translated = GoogleTranslator(source='auto', target=lang).translate(text)
            self.text_edit.setText(translated)
        except Exception as e:
            QMessageBox.critical(self, "Translation Error", str(e))

    # Voice Typing
    def start_voice_typing(self):
        if not self.voice_thread or not self.voice_thread.isRunning():
            self.voice_thread = VoiceTypingThread(self.model_path)
            self.voice_thread.text_output.connect(self.append_text)
            self.voice_thread.start()

    def stop_voice_typing(self):
        if self.voice_thread and self.voice_thread.isRunning():
            self.voice_thread.stop()

    def append_text(self, text):
        current = self.text_edit.toPlainText()
        self.text_edit.setText(current + " " + text)

    # Password & Encryption
    def set_password(self):
        password, ok = QInputDialog.getText(self, "Set Password", "Enter a new password:")
        if ok and password:
            self.current_password = password
            self.recovery_key = base64.urlsafe_b64encode(secrets.token_bytes(16)).decode()
            QMessageBox.information(self, "Password Set", f"Recovery Key: {self.recovery_key}")

    def save_encrypted_file(self):
        if not self.current_password:
            QMessageBox.warning(self, "Password Required", "Set a password first.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Save Encrypted File", "", "Encrypted (*.enc)")
        if path:
            data = self.text_edit.toPlainText().encode()
            encrypted = encrypt_data(data, self.current_password)
            with open(path, 'wb') as f:
                f.write(encrypted)

    def open_encrypted_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Encrypted File", "", "Encrypted (*.enc)")
        if path:
            password, ok = QInputDialog.getText(self, "Enter Password", "Enter password to decrypt:")
            if ok and password:
                with open(path, 'rb') as f:
                    encrypted = f.read()
                    try:
                        decrypted = decrypt_data(encrypted, password)
                        self.text_edit.setText(decrypted.decode())
                        self.current_password = password
                    except Exception:
                        QMessageBox.critical(self, "Error", "Failed to decrypt. Invalid password?")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = TextCraftAI()
    window.show()
    sys.exit(app.exec())
